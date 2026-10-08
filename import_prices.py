"""
Michelin Amazon mastersheet -> asin_prices.csv (then load with load_local.py).

    python3 -m pip install openpyxl            # once
    python3 import_prices.py MICHELIN_AMAZON_MASTERSHEET.xlsx
    python3 import_prices.py SHEET.xlsx --valid-from 2026-11-01   # a price change

Reads the cost price ("ETrade CP W/o GST" = what Amazon pays per unit, ex GST),
Selling Price and MRP per ASIN. Tabs are read in priority order, first numeric
CP wins:
    CONVERTED LISTINGS  (listings changed to a new product - newest truth)
    ASIN MAPPING        (the maintained mapping)
    ALL AMZ SKU         (fallback; has some stale prices, e.g. B07THV6XJ6)

--valid-from: the date these prices apply from. The first import uses
2000-01-01, i.e. today's prices value all history (agreed Oct 2026). For a
later price change, import the new sheet with its effective date: the new
rows are ADDED to asin_prices.csv and older sales keep their old price.
"""

import argparse
import csv
import os
import sys

TABS = ["CONVERTED LISTINGS", "ASIN MAPPING", "ALL AMZ SKU"]
OUT = "asin_prices.csv"
FIELDS = ["asin", "valid_from", "cost_price_ex_gst", "selling_price", "mrp", "india_sku", "source_tab"]


def num(v):
    return round(float(v), 2) if isinstance(v, (int, float)) else None


def read_sheet(path):
    try:
        import openpyxl
    except ImportError:
        sys.exit("Run: python3 -m pip install openpyxl")
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    prices = {}
    for tab in TABS:
        if tab not in wb.sheetnames:
            print(f"  tab not found, skipped: {tab}")
            continue
        rows = wb[tab].iter_rows(values_only=True)
        header = [str(h).strip() if h else "" for h in next(rows)]
        n = 0
        for r in rows:
            d = dict(zip(header, r))
            asin = str(d.get("Amazon ASIN") or "").strip().upper()
            cp = num(d.get("ETrade CP W/o GST"))
            if len(asin) != 10 or not asin.startswith("B") or asin in prices or cp is None:
                continue
            sku = d.get("India SKU")
            prices[asin] = {"asin": asin, "cost_price_ex_gst": cp,
                            "selling_price": num(d.get("Selling Price")), "mrp": num(d.get("MRP")),
                            "india_sku": str(sku).removesuffix(".0").strip() if sku is not None else "",
                            "source_tab": tab}
            n += 1
        print(f"  {tab}: {n} ASIN prices")
    return prices


if __name__ == "__main__":
    p = argparse.ArgumentParser()
    p.add_argument("xlsx")
    p.add_argument("--valid-from", default="2000-01-01")
    args = p.parse_args()

    new = read_sheet(args.xlsx)
    for row in new.values():
        row["valid_from"] = args.valid_from

    kept = []
    if os.path.exists(OUT):
        with open(OUT, newline="", encoding="utf-8-sig") as f:
            # Replace rows with the same valid_from (re-import), keep other dates.
            kept = [r for r in csv.DictReader(f) if r["valid_from"] != args.valid_from]
    rows = kept + sorted(new.values(), key=lambda r: r["asin"])
    with open(OUT, "w", newline="", encoding="utf-8-sig") as f:
        w = csv.DictWriter(f, fieldnames=FIELDS)
        w.writeheader()
        w.writerows(rows)
    print(f"Wrote {OUT}: {len(new)} prices valid from {args.valid_from} "
          f"({len(kept)} rows from other dates kept). Now run: python3 load_local.py")
